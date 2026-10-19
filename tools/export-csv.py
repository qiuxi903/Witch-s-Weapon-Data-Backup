# 把五语对照剧情导出为 CSV / Excel
import io, os, sys, re, csv, json

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

SRC = r'G:\project\其他\game\剧情文本_解密版\国服v2.0.1'
EXTRA = r'G:\project\其他\game\剧情文本_解密版\怀旧服v128_未上线剧情'
OUT = r'G:\project\其他\game\剧情文本_表格版'
os.makedirs(OUT, exist_ok=True)

LANG = ['简体中文', '繁體中文', '日本語', '한국어', 'English', '语言6', '语言7', '语言8']

# Excel 不允许控制字符（除 \t \n \r）与非字符
_ILLEGAL = re.compile(
    r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f'
    r'\ud800-\udfff\ufffe\uffff'
    r'\U0001FFFE\U0001FFFF\U0010FFFE\U0010FFFF]'
)
def clean(s):
    return _ILLEGAL.sub('', s)

def chapter_of(num):
    if num < 10000: return '序章·系统与引导'
    n = num // 10000
    if n == 1: return '第一章·人生的终结、命运的开始'
    if n == 2: return '第二章·想守护的她、被守护的你'
    if n == 3:
        return '活动剧情' if num >= 31000 else '第三章·神魂的治愈、身体的诅咒'
    if n == 4: return '支线剧情'
    return '测试与其它'

rows = []
for src, tag in [(SRC, ''), (EXTRA, '【未上线】')]:
    if not os.path.isdir(src):
        continue
    for f in sorted(os.listdir(src)):
        m = re.match(r'sentence_(\d+)\.txt$', f)
        if not m:
            continue
        num = int(m.group(1))
        c = io.open(os.path.join(src, f), encoding='utf-8').read()
        for ln in c.split('\n'):
            if not ln.strip():
                continue
            p = ln.split('\t')
            if len(p) < 3:
                continue
            row = {
                '文件': tag + f[:-4],
                '编号': num,
                '章节': chapter_of(num),
                '行号': clean(p[0].strip()),
                '说话人': clean(p[1].strip()),
            }
            for i in range(5):
                row[LANG[i]] = clean(p[2 + i].strip()) if len(p) > 2 + i else ''
            rows.append(row)

print('总行数: %d' % len(rows))

# ---- CSV ----
cols = ['文件', '编号', '章节', '行号', '说话人'] + LANG[:5]
csvp = os.path.join(OUT, '魔女兵器_剧情_五语对照.csv')
with io.open(csvp, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore')
    w.writeheader()
    w.writerows(rows)
print('CSV: %s (%.2f MB)' % (csvp, os.path.getsize(csvp) / 1024 / 1024))

# ---- 按章节分表 ----
from collections import OrderedDict
bychap = OrderedDict()
for r in rows:
    bychap.setdefault(r['章节'], []).append(r)
for chap, rs in bychap.items():
    safe = re.sub(r'[\\/:*?"<>|]', '_', chap)
    p = os.path.join(OUT, '分章节_%s.csv' % safe)
    with io.open(p, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore')
        w.writeheader()
        w.writerows(rs)
    print('  %-40s %6d 行' % (chap, len(rs)))

# ---- XLSX（多工作表）----
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    head_fill = PatternFill('solid', fgColor='1F3864')
    head_font = Font(color='FFFFFF', bold=True, size=10)
    widths = [16, 8, 22, 6, 14, 60, 60, 60, 60, 60]

    def add_sheet(title, rs):
        ws = wb.create_sheet(title[:31])
        ws.append(cols)
        for i, cl in enumerate(cols, 1):
            c = ws.cell(row=1, column=i)
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(vertical='center', horizontal='center')
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        for r in rs:
            ws.append([r.get(k, '') for k in cols])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), len(rs) + 1)
        return ws

    add_sheet('全部', rows)
    for chap, rs in bychap.items():
        add_sheet(re.sub(r'[\\/:*?"<>|]', '_', chap), rs)

    xp = os.path.join(OUT, '魔女兵器_剧情_五语对照.xlsx')
    wb.save(xp)
    print('XLSX: %s (%.2f MB, %d 个工作表)' % (xp, os.path.getsize(xp) / 1024 / 1024, len(wb.sheetnames)))
except ImportError as e:
    print('openpyxl 不可用: %s' % e)
